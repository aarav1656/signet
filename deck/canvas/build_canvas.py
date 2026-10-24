# uv run --with python-pptx python build_canvas.py
import json, copy, pathlib
from pptx import Presentation

HERE = pathlib.Path(__file__).parent
TPL = HERE.parents[1].parent.parent.parent / ".progress" / "templates"
KEEP = 4  # zero-based index of slide 5, the blank template


def set_para(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def build(spec):
    prs = Presentation(TPL / spec["template"])
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for i, sid in enumerate(ids):
        if i != KEEP:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
    slide = prs.slides[0]
    done = set()
    for sh in slide.shapes:
        if sh.shape_id == 529:
            for p, t in zip(sh.text_frame.paragraphs, spec["header"]):
                set_para(p, t)
            done.add(529)
        elif sh.shape_id in (525, 526):
            # template year and duplicate wordmark: Google's own render hides them, LibreOffice draws them
            for r in sh.text_frame.paragraphs[0].runs:
                r.text = ""
        elif str(sh.shape_id) in spec["boxes"]:
            set_para(sh.text_frame.paragraphs[1], spec["boxes"][str(sh.shape_id)])
            done.add(sh.shape_id)
    assert done == {529, *map(int, spec["boxes"])}, done
    prs.save(HERE / (spec["file"] + ".pptx"))


# ---- deck/canvas.html: same text as the pptx, one source (content.json) ----
from html import escape

CSS = """/* Hallmark · macrostructure: Workbench (box grid) · tone: technical · anchor hue: amber 80
 * theme: deck/DESIGN.md gate-log tokens, no catalog rotation · nav: none · footer: numbered source strip
 * Two fixed 1920x1080 slides, text generated from deck/canvas/content.json so it equals the official pptx.
 */
:root {
  --color-bg: #121315; --color-panel: #17191c; --color-line: #26292e;
  --color-ink: #e9e7e2; --color-soft: #b5b3ad; --color-mute: #8d9097;
  --color-amber: #f2b134; --color-amber-wash: rgba(242, 177, 52, 0.1);
  --color-amber-edge: rgba(242, 177, 52, 0.4); --color-glow: rgba(242, 177, 52, 0.07);
  --font-display: "Archivo", "Helvetica Neue", Arial, sans-serif;
  --font-mono: "IBM Plex Mono", ui-monospace, "SFMono-Regular", Menlo, monospace;
  --radius: 6px;
}
*, *::before, *::after { box-sizing: border-box; }
html, body { margin: 0; overflow-x: clip; background: var(--color-bg); color: var(--color-ink); }
body { font-family: var(--font-display); line-height: 1.3; }
@page { size: 1920px 1080px; margin: 0; }
.slide { width: 1920px; height: 1080px; padding: 28px 48px 0; display: grid; grid-template-rows: auto 1fr auto; gap: 14px;
  background: var(--color-bg) radial-gradient(900px 500px at 88% -6%, var(--color-glow), transparent 70%);
  position: relative; overflow: hidden; break-after: page; transform-origin: top left; }
.slide + .slide { margin-top: 24px; }
.head { display: flex; align-items: baseline; gap: 28px; }
.head h1 { margin: 0; font-size: 34px; font-weight: 600; font-stretch: 96%; letter-spacing: -0.035em; line-height: 1; font-style: normal; }
.head .files { margin-left: auto; font-family: var(--font-mono); font-size: 16px; color: var(--color-mute); }
.head .files a { color: var(--color-amber); margin-left: 18px; text-decoration: underline; text-underline-offset: 4px; }
.grid { display: grid; grid-template-columns: repeat(12, minmax(0, 1fr)); grid-template-rows: auto repeat(3, 1fr); gap: 14px; min-height: 0; }
.box { background: var(--color-panel); border: 1px solid var(--color-line); border-radius: var(--radius); padding: 16px 24px 12px; min-width: 0; min-height: 0; overflow: hidden; display: flex; flex-direction: column; grid-column: span 4; }
.box.wide { grid-column: span 6; }
.box.top { grid-column: span 12; }
.box.key { border-color: var(--color-amber-edge); background: var(--color-amber-wash); }
.lbl { margin: 0 0 12px; font-family: var(--font-mono); font-size: 17px; font-weight: 400; letter-spacing: 0.1em; text-transform: uppercase; color: var(--color-mute); }
.txt { margin: 0; font-size: 22px; line-height: 1.28; color: var(--color-soft); overflow-wrap: anywhere; }
.wide .txt { font-size: 24px; }
.top .txt { font-size: 32px; color: var(--color-ink); font-weight: 600; letter-spacing: -0.02em; }
.src { margin: 0; padding: 12px 0 0; border-top: 1px solid var(--color-line); font-family: var(--font-mono); font-size: 12.5px; line-height: 1.4; color: var(--color-mute); columns: 3; column-gap: 32px; margin-bottom: 14px; list-style: none; }
.src li { display: block; break-inside: avoid; }
.src b { color: var(--color-amber); font-weight: 500; }
@media print { html, body { background: var(--color-bg); } .slide { transform: none !important; margin: 0 !important; } .head .files { display: none; } }
"""

SCRIPT = """(function () {
  if (matchMedia('print').matches) return;
  function fit() {
    var k = Math.min(1, innerWidth / 1920);
    document.querySelectorAll('.slide').forEach(function (s) {
      s.style.transform = 'scale(' + k + ')';
      s.style.marginBottom = (1080 * k - 1080) + 'px';
      s.style.marginRight = (1920 * k - 1920) + 'px';
    });
  }
  addEventListener('resize', fit); fit();
})();"""

ORDER = [530, 531, 534, 532, 533, 535, 536, 537]  # reading order of the grid, top row then middle then bottom
KEY = {532, 537}  # the wedge and the weekend ship carry the amber edge, as in the template's highlighted wedge box


def titles(spec):
    prs = Presentation(TPL / spec["template"])
    return {sh.shape_id: sh.text_frame.paragraphs[0].text for sh in prs.slides[KEEP].shapes if sh.shape_id in ORDER}


def section(ident, label, spec):
    t = titles(spec)
    boxes = [f'    <div class="box top key"><p class="lbl">{escape(spec["header"][0])}</p><p class="txt">{escape(spec["header"][1])}</p></div>']
    for sid in ORDER:
        cls = "box" + (" wide" if sid in (536, 537) else "") + (" key" if sid in KEY else "")
        boxes.append(f'    <div class="{cls}"><p class="lbl">{escape(t[sid])}</p><p class="txt">{escape(spec["boxes"][str(sid)])}</p></div>')
    src = "\n".join(f'    <li><b>{i}</b> {escape(s)}</li>' for i, s in enumerate(spec["sources"], 1))
    f = spec["file"]
    return f"""<section class="slide" id="{ident}" aria-label="{label}">
  <div class="head"><h1>{label}</h1><span class="files">Official file <a href="canvas/{f}.pptx">pptx</a><a href="canvas/{f}.pdf">pdf</a><a href="canvas/{f}.png">png</a></span></div>
  <div class="grid">
{chr(10).join(boxes)}
  </div>
  <ol class="src">
{src}
  </ol>
</section>"""


def write_html(c):
    body = section("product", "Product canvas", c["product"]) + "\n\n" + section("gtm", "GTM canvas", c["gtm"])
    page = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Signet: Product canvas and GTM canvas</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Archivo:wdth,wght@100..125,400..700&family=IBM+Plex+Mono:wght@400;500&display=swap" rel="stylesheet">
<style>
{CSS}</style>
</head>
<body>
{body}
<script>
{SCRIPT}
</script>
</body>
</html>
"""
    (HERE.parent / "canvas.html").write_text(page)


if __name__ == "__main__":
    c = json.loads((HERE / "content.json").read_text())
    for k in ("product", "gtm"):
        build(c[k])
    write_html(c)
